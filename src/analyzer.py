from pathlib import Path
from openpyxl import load_workbook

from type_inference import infer_column_type


def normalize_value(value):
    if value is None:
        return ""
    return str(value).strip()


def row_values(ws, row_number):
    return [
        normalize_value(cell.value)
        for cell in ws[row_number]
    ]


def detect_header(ws):
    max_scan = min(ws.max_row, 30)
    candidates = []

    for row_number in range(1, max_scan + 1):
        values = row_values(ws, row_number)
        non_empty = [v for v in values if v]

        if len(non_empty) < 2:
            continue

        if row_number < ws.max_row:
            next_values = row_values(ws, row_number + 1)
            if not any(next_values):
                continue

        candidates.append((row_number, len(non_empty)))

    if not candidates:
        return None

    candidates.sort(key=lambda x: (-x[1], x[0]))

    return candidates[0][0]


def analyze_sheet(ws):
    print()
    print("=" * 60)
    print(f"Sheet: {ws.title}")
    print(f"尺寸: {ws.max_row} 行 × {ws.max_column} 列")
    print("=" * 60)

    header_row = detect_header(ws)

    if header_row is None:
        print("❌ 未找到可靠表头")
        return

    print(f"✅ 表头行: 第 {header_row} 行")
    print()
    print("字段类型:")

    for column_index in range(1, ws.max_column + 1):

        header = ws.cell(
            header_row,
            column_index
        ).value

        if header is None:
            continue

        header = str(header).strip()

        values = []

        for row_number in range(
            header_row + 1,
            ws.max_row + 1
        ):
            values.append(
                ws.cell(
                    row_number,
                    column_index
                ).value
            )

        pg_type = infer_column_type(values)

        print(
            f"  {column_index}. "
            f"{header} → {pg_type}"
        )


def analyze_excel(file_path):
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"文件不存在: {path}"
        )

    if path.suffix.lower() not in {
        ".xlsx",
        ".xlsm"
    }:
        raise ValueError(
            "目前只支持 .xlsx / .xlsm 文件"
        )

    wb = load_workbook(
        filename=path,
        read_only=False,
        data_only=False,
    )

    print("=" * 60)
    print(f"Excel: {path.name}")
    print(f"Sheet 数量: {len(wb.worksheets)}")
    print("=" * 60)

    for ws in wb.worksheets:
        analyze_sheet(ws)

    print()
    print("=" * 60)
    print("类型分析完成")
    print("=" * 60)


if __name__ == "__main__":
    import sys

    if len(sys.argv) != 2:
        print(
            "用法: "
            "python src/analyzer.py <Excel文件>"
        )
        sys.exit(1)

    analyze_excel(sys.argv[1])
