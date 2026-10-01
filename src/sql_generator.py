from type_inference import infer_column_type


def quote_identifier(name):
    """
    PostgreSQL 安全标识符。
    中文、空格、特殊字符都可以保留。
    """
    name = str(name).strip()

    if not name:
        name = "unnamed_column"

    return '"' + name.replace('"', '""') + '"'


def make_unique_names(headers):
    """
    处理重复列名。
    例如：
    商品 → 商品
    商品 → 商品_2
    """
    used = {}
    result = []

    for header in headers:
        name = str(header).strip()

        if not name:
            name = "unnamed_column"

        if name not in used:
            used[name] = 1
            result.append(name)
        else:
            used[name] += 1
            result.append(
                f"{name}_{used[name]}"
            )

    return result


def generate_create_table(ws, header_row):
    headers = [
        ws.cell(
            header_row,
            column_index
        ).value
        for column_index in range(
            1,
            ws.max_column + 1
        )
    ]

    headers = [
        "" if value is None else str(value).strip()
        for value in headers
    ]

    headers = make_unique_names(headers)

    columns = []

    for column_index, header in enumerate(
        headers,
        start=1
    ):
        values = []

        for row_number in range(
            header_row + 1,
            ws.max_row + 1
        ):
            values.append(
                ws.cell(
                    row_number,
                    column_index
                ).value
            )

        pg_type = infer_column_type(values)

        columns.append(
            f"    {quote_identifier(header)} "
            f"{pg_type}"
        )

    table_name = quote_identifier(ws.title)

    sql = (
        f"CREATE TABLE {table_name} (\n"
        + ",\n".join(columns)
        + "\n);"
    )

    return sql


if __name__ == "__main__":
    import sys
    from pathlib import Path
    from openpyxl import load_workbook

    if len(sys.argv) != 2:
        print(
            "用法: "
            "python src/sql_generator.py <Excel文件>"
        )
        sys.exit(1)

    file_path = Path(sys.argv[1])

    wb = load_workbook(
        filename=file_path,
        read_only=False,
        data_only=False,
    )

    for ws in wb.worksheets:

        # 简单使用前 30 行寻找表头
        header_row = None

        for row_number in range(
            1,
            min(ws.max_row, 30) + 1
        ):
            values = [
                ws.cell(
                    row_number,
                    column_index
                ).value
                for column_index in range(
                    1,
                    ws.max_column + 1
                )
            ]

            if sum(
                value is not None
                and str(value).strip() != ""
                for value in values
            ) >= 2:
                header_row = row_number
                break

        if header_row is None:
            print(
                f"\n❌ {ws.title}: "
                "无法找到表头"
            )
            continue

        print()
        print("=" * 60)
        print(f"Sheet: {ws.title}")
        print("=" * 60)
        print()
        print(
            generate_create_table(
                ws,
                header_row
            )
        )
