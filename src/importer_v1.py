from pathlib import Path
import csv
import io
import os

from openpyxl import load_workbook
from dotenv import load_dotenv
import psycopg2
from psycopg2 import sql

from type_inference import infer_column_type


def detect_header(ws):
    """寻找前 30 行中最可能的表头。"""
    candidates = []

    for row_number in range(1, min(ws.max_row, 30) + 1):
        values = [
            ws.cell(row_number, col).value
            for col in range(1, ws.max_column + 1)
        ]

        non_empty = [
            v
            for v in values
            if v is not None and str(v).strip()
        ]

        if len(non_empty) >= 2:
            candidates.append(
                (row_number, len(non_empty))
            )

    if not candidates:
        return None

    candidates.sort(
        key=lambda x: (-x[1], x[0])
    )

    return candidates[0][0]


def make_unique_names(headers):
    """保证 PostgreSQL 字段名唯一。"""
    used = {}
    result = []

    for header in headers:
        name = (
            str(header).strip()
            if header is not None
            else ""
        )

        if not name:
            name = "unnamed_column"

        if name not in used:
            used[name] = 1
            result.append(name)
        else:
            used[name] += 1
            result.append(
                f"{name}_{used[name]}"
            )

    return result


def cell_to_text(value, pg_type):
    """
    Excel 单元格 → PostgreSQL COPY 文本。

    boolean 特殊处理：
    是 / 是的 / yes / true / 1 → true
    否 / no / false / 0 → false
    """

    if value is None:
        return ""

    if pg_type == "boolean":
        text = str(value).strip().lower()

        if text in {
            "是",
            "是的",
            "yes",
            "true",
            "1",
        }:
            return "true"

        if text in {
            "否",
            "no",
            "false",
            "0",
        }:
            return "false"

    return str(value)


def import_sheet(ws, conn):
    """导入一个 Sheet。"""

    header_row = detect_header(ws)

    if header_row is None:
        print(
            f"❌ {ws.title}: 找不到表头，跳过"
        )
        return False

    headers = [
        ws.cell(
            header_row,
            col
        ).value
        for col in range(
            1,
            ws.max_column + 1
        )
    ]

    headers = make_unique_names(headers)

    # =========================
    # 判断每一列 PostgreSQL 类型
    # =========================

    column_types = []

    for col in range(
        1,
        ws.max_column + 1
    ):
        values = [
            ws.cell(row, col).value
            for row in range(
                header_row + 1,
                ws.max_row + 1
            )
        ]

        column_types.append(
            infer_column_type(values)
        )

    table = sql.Identifier(ws.title)

    try:
        # =========================
        # 1. 创建表
        # =========================

        columns = []

        for name, pg_type in zip(
            headers,
            column_types
        ):
            columns.append(
                sql.SQL("{} {}").format(
                    sql.Identifier(name),
                    sql.SQL(pg_type),
                )
            )

        create_sql = sql.SQL(
            "CREATE TABLE {} ({})"
        ).format(
            table,
            sql.SQL(", ").join(columns),
        )

        with conn.cursor() as cur:
            cur.execute(create_sql)

        # =========================
        # 2. 准备 COPY 数据
        # =========================

        buffer = io.StringIO()

        writer = csv.writer(
            buffer,
            delimiter="\t",
            lineterminator="\n",
            quoting=csv.QUOTE_MINIMAL,
        )

        data_rows = 0

        for row_number in range(
            header_row + 1,
            ws.max_row + 1
        ):
            values = [
                ws.cell(
                    row_number,
                    col
                ).value
                for col in range(
                    1,
                    ws.max_column + 1
                )
            ]

            # 完全空行跳过
            if not any(
                v is not None
                for v in values
            ):
                continue

            writer.writerow([
                cell_to_text(
                    value,
                    pg_type
                )
                for value, pg_type in zip(
                    values,
                    column_types
                )
            ])

            data_rows += 1

        # =========================
        # 3. COPY 批量导入
        # =========================

        buffer.seek(0)

        copy_sql = sql.SQL(
            """
            COPY {} ({})
            FROM STDIN
            WITH (
                FORMAT CSV,
                DELIMITER E'\\t'
            )
            """
        ).format(
            table,
            sql.SQL(", ").join(
                sql.Identifier(name)
                for name in headers
            ),
        )

        with conn.cursor() as cur:
            cur.copy_expert(
                copy_sql.as_string(conn),
                buffer
            )

        # =========================
        # 4. 验证行数
        # =========================

        with conn.cursor() as cur:
            cur.execute(
                sql.SQL(
                    "SELECT COUNT(*) FROM {}"
                ).format(table)
            )

            actual_count = cur.fetchone()[0]

        # =========================
        # 5. 验证
        # =========================

        if data_rows != actual_count:
            raise RuntimeError(
                f"行数不一致: "
                f"Excel={data_rows}, "
                f"PostgreSQL={actual_count}"
            )

        conn.commit()

        print()
        print(f"✅ Sheet: {ws.title}")
        print(
            f"   表头: 第 {header_row} 行"
        )
        print(
            f"   表名: {ws.title}"
        )
        print(
            f"   字段: {len(headers)}"
        )
        print(
            f"   Excel 数据行: {data_rows}"
        )
        print(
            f"   PostgreSQL 行数: {actual_count}"
        )
        print(
            "   ✅ 行数验证通过"
        )

        print("   字段类型:")

        for name, pg_type in zip(
            headers,
            column_types
        ):
            print(
                f"      {name} → {pg_type}"
            )

        return True

    except Exception:
        # 当前 Sheet 出错时回滚
        conn.rollback()

        print()
        print(
            f"❌ Sheet 导入失败: {ws.title}"
        )

        raise


def main(file_path):
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"文件不存在: {path}"
        )

    if path.suffix.lower() not in {
        ".xlsx",
        ".xlsm"
    }:
        raise ValueError(
            "目前只支持 .xlsx / .xlsm"
        )

    # =========================
    # 读取数据库配置
    # =========================

    load_dotenv()

    required = [
        "PGHOST",
        "PGPORT",
        "PGDATABASE",
        "PGUSER",
        "PGPASSWORD",
    ]

    missing = [
        key
        for key in required
        if not os.getenv(key)
    ]

    if missing:
        raise RuntimeError(
            "缺少数据库配置: "
            + ", ".join(missing)
        )

    # =========================
    # 连接 PostgreSQL
    # =========================

    conn = psycopg2.connect(
        host=os.getenv("PGHOST"),
        port=os.getenv("PGPORT"),
        dbname=os.getenv("PGDATABASE"),
        user=os.getenv("PGUSER"),
        password=os.getenv("PGPASSWORD"),
    )

    print("=" * 60)
    print("Excel → PostgreSQL")
    print("=" * 60)

    print(
        f"Excel: {path.name}"
    )
    print(
        f"数据库: {conn.info.dbname}"
    )
    print(
        f"用户: {conn.info.user}"
    )
    print()

    # =========================
    # 读取 Excel
    # =========================

    wb = load_workbook(
        filename=path,
        read_only=False,

        # ★★★ 关键设置 ★★★
        # 读取公式计算后的值，
        # 不读取 =VLOOKUP(...) / =SUM(...) 等公式
        data_only=True,
    )

    print(
        f"Sheet 数量: {len(wb.worksheets)}"
    )

    success_count = 0

    try:
        for ws in wb.worksheets:

            # =========================
            # 如果表已经存在，停止
            # =========================

            with conn.cursor() as cur:
                cur.execute(
                    """
                    SELECT EXISTS (
                        SELECT 1
                        FROM information_schema.tables
                        WHERE table_schema = 'public'
                        AND table_name = %s
                    )
                    """,
                    (ws.title,),
                )

                exists = cur.fetchone()[0]

            if exists:
                raise RuntimeError(
                    f"表已存在: {ws.title}\n"
                    "为避免覆盖已有数据，程序已停止。"
                )

            if import_sheet(ws, conn):
                success_count += 1

    finally:
        conn.close()

    print()
    print("=" * 60)
    print(
        f"🎉 导入完成: "
        f"{success_count}/"
        f"{len(wb.worksheets)} 个 Sheet"
    )
    print("=" * 60)


if __name__ == "__main__":
    import sys

    if len(sys.argv) != 2:
        print("用法:")
        print(
            "  python src/importer.py <Excel文件>"
        )
        sys.exit(1)

    try:
        main(sys.argv[1])

    except Exception as e:
        print()
        print("=" * 60)
        print("❌ 导入失败")
        print("=" * 60)
        print(e)
        sys.exit(1)
