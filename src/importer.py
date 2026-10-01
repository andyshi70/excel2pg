from pathlib import Path
import csv
import io
import os
import sys

from openpyxl import load_workbook
from dotenv import load_dotenv
import psycopg2
from psycopg2 import sql

from type_inference import infer_column_type

# 要导入的 Excel 统一放这个目录；无参数运行 = 批量导入这里所有文件（WebUI 导入按钮走这条）
EXCEL_DIR = Path("/Users/evandy/excel2pg/src/excel")


def pg_ident(name: str) -> str:
    """PG 标识符上限 63 字节：按字节截断且不切坏 UTF-8（否则 PG 默默截断，前后名字对不上）。"""
    return name.encode("utf-8")[:63].decode("utf-8", "ignore")


def table_name_for(path, ws) -> str:
    """表名 = excel名_sheet名（防止不同 Excel 的同名 Sheet 互相覆盖）。"""
    return pg_ident(f"{Path(path).stem}_{ws.title}")


def detect_header(ws):
    """寻找前 30 行中最可能的表头。"""
    candidates = []

    for row_number in range(1, min(ws.max_row, 30) + 1):
        values = [
            ws.cell(row_number, col).value
            for col in range(1, ws.max_column + 1)
        ]

        non_empty = [
            v
            for v in values
            if v is not None and str(v).strip()
        ]

        if len(non_empty) >= 2:
            candidates.append((row_number, len(non_empty)))

    if not candidates:
        return None

    candidates.sort(key=lambda x: (-x[1], x[0]))
    return candidates[0][0]


def make_unique_names(headers):
    """保证 PostgreSQL 字段名唯一。"""
    used = {}
    result = []

    for header in headers:
        name = str(header).strip() if header is not None else ""

        if not name:
            name = "unnamed_column"

        if name not in used:
            used[name] = 1
            result.append(name)
        else:
            used[name] += 1
            result.append(f"{name}_{used[name]}")

    return result


def cell_to_text(value, pg_type):
    """
    Excel 单元格 → PostgreSQL COPY 文本。

    boolean 特殊处理：
    是 / 是的 / yes / true / 1 → true
    否 / no / false / 0 → false
    """
    if value is None:
        return ""

    if pg_type == "boolean":
        text = str(value).strip().lower()

        if text in {"是", "是的", "yes", "true", "1"}:
            return "true"

        if text in {"否", "no", "false", "0"}:
            return "false"

    return str(value)


def table_exists(conn, table_name):
    """检查 public schema 下表是否已存在。"""
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT EXISTS (
                SELECT 1
                FROM information_schema.tables
                WHERE table_schema = 'public'
                  AND table_name = %s
            )
            """,
            (table_name,),
        )
        return cur.fetchone()[0]


def drop_table(conn, table_name):
    """删除已存在的表（用于重新导入）。"""
    with conn.cursor() as cur:
        cur.execute(
            sql.SQL("DROP TABLE IF EXISTS {} CASCADE").format(
                sql.Identifier(table_name)
            )
        )
    conn.commit()


def prompt_update_choice(existing_sheets):
    """
    交互式让用户选择要更新的已存在表。
    返回需要更新的 sheet 标题集合。
    """
    print()
    print("=" * 60)
    print("以下 Sheet 对应的表已存在于数据库中：")
    print("=" * 60)

    for i, name in enumerate(existing_sheets, 1):
        print(f"  {i}. {name}")

    print()
    print("请选择要【更新/重新导入】的表（旧数据会被覆盖）：")
    print("  - 输入序号，多个用逗号分隔，例如：1,3")
    print("  - 输入 all 更新全部已存在的表")
    print("  - 输入 skip 或直接回车：全部跳过，不更新已存在的表")
    print()

    while True:
        raw = input("你的选择: ").strip().lower()

        if raw == "" or raw == "skip":
            print("→ 已选择：跳过所有已存在的表")
            return set()

        if raw == "all":
            print(f"→ 已选择：更新全部 {len(existing_sheets)} 个表")
            return set(existing_sheets)

        try:
            indices = [int(x.strip()) for x in raw.split(",") if x.strip()]
            chosen = set()
            invalid = []

            for idx in indices:
                if 1 <= idx <= len(existing_sheets):
                    chosen.add(existing_sheets[idx - 1])
                else:
                    invalid.append(str(idx))

            if invalid:
                print(f"  无效序号: {', '.join(invalid)}，请重新输入")
                continue

            if not chosen:
                print("  未选中任何表，请重新输入")
                continue

            print(f"→ 已选择更新: {', '.join(chosen)}")
            return chosen

        except ValueError:
            print("  输入格式不对，请按提示重新输入")


def import_sheet(ws, conn, *, replace=False, table_name=None):
    """
    导入一个 Sheet。table_name = 库里表名（excel名_sheet名）；缺省回退 sheet 名。
    replace=True 时，若表已存在会先 DROP 再重建。
    """
    table_name = table_name or ws.title
    header_row = detect_header(ws)

    if header_row is None:
        print(f"❌ {ws.title}: 找不到表头，跳过")
        return False

    headers = [
        ws.cell(header_row, col).value
        for col in range(1, ws.max_column + 1)
    ]
    headers = make_unique_names(headers)

    # 推断每列 PostgreSQL 类型
    column_types = []
    for col in range(1, ws.max_column + 1):
        values = [
            ws.cell(row, col).value
            for row in range(header_row + 1, ws.max_row + 1)
        ]
        column_types.append(infer_column_type(values))

    table = sql.Identifier(table_name)

    try:
        # 需要替换时先删表
        if replace:
            print(f"   ⟳ 正在删除旧表: {table_name}")
            drop_table(conn, table_name)

        # 1. 创建表
        columns = []
        for name, pg_type in zip(headers, column_types):
            columns.append(
                sql.SQL("{} {}").format(
                    sql.Identifier(name),
                    sql.SQL(pg_type),
                )
            )

        create_sql = sql.SQL("CREATE TABLE {} ({})").format(
            table,
            sql.SQL(", ").join(columns),
        )

        with conn.cursor() as cur:
            cur.execute(create_sql)

        # 2. 准备 COPY 数据
        buffer = io.StringIO()
        writer = csv.writer(
            buffer,
            delimiter="\t",
            lineterminator="\n",
            quoting=csv.QUOTE_MINIMAL,
        )

        data_rows = 0
        for row_number in range(header_row + 1, ws.max_row + 1):
            values = [
                ws.cell(row_number, col).value
                for col in range(1, ws.max_column + 1)
            ]

            # 完全空行跳过
            if not any(v is not None for v in values):
                continue

            writer.writerow([
                cell_to_text(value, pg_type)
                for value, pg_type in zip(values, column_types)
            ])
            data_rows += 1

        # 3. COPY 批量导入
        buffer.seek(0)
        copy_sql = sql.SQL(
            """
            COPY {} ({})
            FROM STDIN
            WITH (
                FORMAT CSV,
                DELIMITER E'\\t'
            )
            """
        ).format(
            table,
            sql.SQL(", ").join(sql.Identifier(name) for name in headers),
        )

        with conn.cursor() as cur:
            cur.copy_expert(copy_sql.as_string(conn), buffer)

        # 4. 验证行数
        with conn.cursor() as cur:
            cur.execute(
                sql.SQL("SELECT COUNT(*) FROM {}").format(table)
            )
            actual_count = cur.fetchone()[0]

        if data_rows != actual_count:
            raise RuntimeError(
                f"行数不一致: Excel={data_rows}, PostgreSQL={actual_count}"
            )

        conn.commit()

        action = "更新" if replace else "新建"
        print()
        print(f"✅ Sheet: {ws.title}（{action}）")
        print(f"   表头: 第 {header_row} 行")
        print(f"   表名: {table_name}")
        print(f"   字段: {len(headers)}")
        print(f"   Excel 数据行: {data_rows}")
        print(f"   PostgreSQL 行数: {actual_count}")
        print("   ✅ 行数验证通过")
        print("   字段类型:")
        for name, pg_type in zip(headers, column_types):
            print(f"      {name} → {pg_type}")

        return True

    except Exception:
        conn.rollback()
        print()
        print(f"❌ Sheet 导入失败: {ws.title}")
        raise


def main(file_path, update_existing=None):
    """
    update_existing:
      None  = 交互式询问哪些已存在表要覆盖（原行为）
      True  = 已存在的表全部覆盖重导
      False = 已存在的表全部跳过
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"文件不存在: {path}")

    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("目前只支持 .xlsx / .xlsm")

    load_dotenv()

    required = ["PGHOST", "PGPORT", "PGDATABASE", "PGUSER", "PGPASSWORD"]
    missing = [key for key in required if not os.getenv(key)]
    if missing:
        raise RuntimeError("缺少数据库配置: " + ", ".join(missing))

    conn = psycopg2.connect(
        host=os.getenv("PGHOST"),
        port=os.getenv("PGPORT"),
        dbname=os.getenv("PGDATABASE"),
        user=os.getenv("PGUSER"),
        password=os.getenv("PGPASSWORD"),
    )

    print("=" * 60)
    print("Excel → PostgreSQL")
    print("=" * 60)
    print(f"Excel: {path.name}")
    print(f"数据库: {conn.info.dbname}")
    print(f"用户: {conn.info.user}")
    print()

    # data_only=True：读取公式计算后的值，不读公式本身
    wb = load_workbook(
        filename=path,
        read_only=False,
        data_only=True,
    )

    print(f"Sheet 数量: {len(wb.worksheets)}")

    # ---------- 先扫描：哪些表已存在 ----------
    existing_sheets = []
    new_sheets = []

    for ws in wb.worksheets:
        tname = table_name_for(path, ws)
        if table_exists(conn, tname):
            existing_sheets.append(tname)
        else:
            new_sheets.append(tname)

    if existing_sheets:
        print()
        print(f"已存在表: {len(existing_sheets)} 个 → {', '.join(existing_sheets)}")
    if new_sheets:
        print(f"新表:     {len(new_sheets)} 个 → {', '.join(new_sheets)}")

    # ---------- 让用户选择要更新哪些已存在的表 ----------
    to_update = set()
    if existing_sheets:
        if update_existing is None:
            to_update = prompt_update_choice(existing_sheets)
        elif update_existing:
            to_update = set(existing_sheets)
            print(f"→ 批量模式：覆盖重导全部已存在表（{len(to_update)} 个）")
        else:
            print("→ 批量模式：跳过已存在表")

    # ---------- 按选择执行导入 ----------
    success_count = 0
    skipped = []

    try:
        for ws in wb.worksheets:
            tname = table_name_for(path, ws)
            exists = tname in existing_sheets

            if exists and tname not in to_update:
                print()
                print(f"⏭  跳过（表已存在且未选择更新）: {tname}")
                skipped.append(tname)
                continue

            replace = exists and tname in to_update
            if import_sheet(ws, conn, replace=replace, table_name=tname):
                success_count += 1

    finally:
        conn.close()

    print()
    print("=" * 60)
    total = len(wb.worksheets)
    print(f"🎉 完成: 成功 {success_count}/{total} 个 Sheet")
    if skipped:
        print(f"   跳过: {len(skipped)} 个 → {', '.join(skipped)}")
    print("=" * 60)


if __name__ == "__main__":
    if len(sys.argv) == 2:
        # 单文件模式（保留原行为，可交互询问覆盖）
        try:
            main(sys.argv[1])
        except Exception as e:
            print()
            print("=" * 60)
            print("❌ 导入失败")
            print("=" * 60)
            print(e)
            sys.exit(1)
    else:
        # 无参数 = 批量导入 EXCEL_DIR 下所有 Excel，已存在表全部覆盖重导
        print("=" * 60)
        print(f"批量导入目录: {EXCEL_DIR}")
        files = sorted(list(EXCEL_DIR.glob("*.xlsx")) + list(EXCEL_DIR.glob("*.xlsm")))
        if not files:
            print(f"目录里没有 Excel 文件，把文件放进这个目录再运行")
            sys.exit(0)
        failed = []
        for f in files:
            print()
            print("#" * 60)
            print(f"文件: {f.name}（{files.index(f) + 1}/{len(files)}）")
            try:
                main(f, update_existing=True)
            except Exception as e:
                print(f"❌ 文件导入失败: {f.name} → {e}")
                failed.append(f.name)
        print()
        print("=" * 60)
        print(f"批量导入结束: 成功 {len(files) - len(failed)}/{len(files)}")
        if failed:
            print(f"失败: {', '.join(failed)}")
        print("=" * 60)
        if failed:
            sys.exit(1)
