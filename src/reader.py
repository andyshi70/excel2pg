from pathlib import Path

from openpyxl import load_workbook


def inspect_excel(file_path: str):
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"文件不存在: {path}")

    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("目前只支持 .xlsx / .xlsm 文件")

    wb = load_workbook(
        filename=path,
        read_only=False,
        data_only=False,
    )

    print("=" * 60)
    print(f"Excel: {path.name}")
    print(f"Sheet 数量: {len(wb.worksheets)}")
    print("=" * 60)

    for ws in wb.worksheets:
        print()
        print(f"【Sheet】{ws.title}")
        print(f"尺寸: {ws.max_row} 行 × {ws.max_column} 列")

        # 合并单元格
        merged_ranges = list(ws.merged_cells.ranges)

        print(f"合并区域: {len(merged_ranges)}")

        if merged_ranges:
            for merged in merged_ranges[:20]:
                print(f"  - {merged}")

            if len(merged_ranges) > 20:
                print(f"  ... 还有 {len(merged_ranges) - 20} 个")

        # 隐藏行
        hidden_rows = [
            row_index
            for row_index, dimension in ws.row_dimensions.items()
            if dimension.hidden
        ]

        print(f"隐藏行: {len(hidden_rows)}")

        if hidden_rows:
            print(f"  {hidden_rows[:30]}")

        # 隐藏列
        hidden_columns = [
            column_letter
            for column_letter, dimension in ws.column_dimensions.items()
            if dimension.hidden
        ]

        print(f"隐藏列: {len(hidden_columns)}")

        if hidden_columns:
            print(f"  {hidden_columns[:30]}")

        # 打印前 10 行
        print("前 10 行数据:")

        for row in ws.iter_rows(
            min_row=1,
            max_row=min(ws.max_row, 10),
            values_only=True,
        ):
            values = list(row)

            # 去掉末尾连续空单元格，方便查看
            while values and values[-1] is None:
                values.pop()

            print(f"  {values}")

    print()
    print("=" * 60)
    print("扫描完成")
    print("=" * 60)


if __name__ == "__main__":
    import sys

    if len(sys.argv) != 2:
        print("用法:")
        print("  python src/reader.py <Excel文件>")
        sys.exit(1)

    inspect_excel(sys.argv[1])
